"""
Reads product_data/PRODUCT_DATA.xlsx and updates the PRODUCT_DB
constant inside static/index.html.
Run via: "UPDATE PRODUCTS - Double Click Me.bat"
"""
import sys, os, re, json

try:
    from openpyxl import load_workbook
except ImportError:
    os.system('python -m pip install openpyxl --quiet')
    from openpyxl import load_workbook

EXCEL_PATH  = os.path.join(os.path.dirname(__file__), 'product_data', 'PRODUCT_DATA.xlsx')
HTML_PATH   = os.path.join(os.path.dirname(__file__), 'static', 'index.html')

def strip_pack(raw):
    import re
    p = re.sub(r'\s+\d+(?:X\d+)?(?:TAB|CAP)$', '', raw, flags=re.IGNORECASE).strip().upper()
    p = re.sub(r'\s+\d+X\d+GM$',               '', p,   flags=re.IGNORECASE).strip()
    p = re.sub(r'\s+\d+\*\d+$',                '', p,   flags=re.IGNORECASE).strip()
    return p

def clean(v):
    if v is None: return ''
    s = str(v).strip()
    return '' if s.upper() in ('-BLANK-', 'BLANK', 'N/A', 'NA') else s.upper()

print(f"  Reading {EXCEL_PATH} ...")
wb = load_workbook(EXCEL_PATH)
ws = wb.active

product_db = {}
skipped = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    if len(row) < 1: continue
    brand, generic, manufacturer, hsn = (list(row) + ['','','',''])[:4]
    brand = clean(brand)
    if not brand:
        skipped += 1
        continue
    key = strip_pack(brand)
    g, m, h = clean(generic), clean(manufacturer), clean(hsn)
    if key in product_db:
        ex = product_db[key]
        if not ex['g'] and g: ex['g'] = g
        if not ex['m'] and m: ex['m'] = m
        if not ex['h'] and h: ex['h'] = h
    else:
        product_db[key] = {'g': g, 'm': m, 'h': h}

print(f"  Loaded {len(product_db)} unique products ({skipped} blank rows skipped)")

js_const = f"const PRODUCT_DB = {json.dumps(product_db, ensure_ascii=False, separators=(',',':'))};"

print(f"  Updating {HTML_PATH} ...")
content = open(HTML_PATH, encoding='utf-8').read()

# Replace existing PRODUCT_DB constant
new_content = re.sub(
    r'const PRODUCT_DB = \{[^;]+\};',
    js_const,
    content,
    count=1
)

if new_content == content:
    print("  [ERROR] Could not find PRODUCT_DB in index.html — file may be corrupted.")
    sys.exit(1)

open(HTML_PATH, 'w', encoding='utf-8').write(new_content)
print(f"  Done! {len(product_db)} products now in the app.")
